"""Self-check for check_deck.py: python3 scripts/test_check_deck.py"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(__file__))
from check_deck import check
from pptx import Presentation
from pptx.util import Pt, Inches


def deck(body_pt, named_title=True):
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1)); t.name = "title" if named_title else "box"
    r = t.text_frame.paragraphs[0].add_run(); r.text = "Title"; r.font.size = Pt(36)
    b = s.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(3))
    r = b.text_frame.paragraphs[0].add_run(); r.text = "Body text"
    if body_pt: r.font.size = Pt(body_pt)
    f = s.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.4)); f.name = "footer source"
    r = f.text_frame.paragraphs[0].add_run(); r.text = "Source: case"; r.font.size = Pt(12)
    path = os.path.join(tempfile.mkdtemp(), "d.pptx"); p.save(path); return path


assert check(deck(20), 90)[0] == []
assert any("16pt" in e for e in check(deck(16), 90)[0])
assert any("no explicit font size" in e for e in check(deck(None), 90)[0])
assert any("no title" in e for e in check(deck(20, named_title=False), 90)[0])
print("check_deck self-check ok")
