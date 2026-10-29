"""Rule checks on a built .pptx. Exit code 1 if anything fails.

Usage: python3 scripts/check_deck.py deck.pptx [--max-words 90]

Checks, per slide:
  - every text run has an explicit font size (inherited sizes can't be verified)
  - body text >= 18pt; shapes named "footer*" (source line, slide number) >= 12pt
  - text boxes probably overflowing (rough estimate from text length, size and box area)
  - the slide has a title (a shape named "title*" or a title placeholder)
  - word count over --max-words (crowding warning only)
"""
import argparse, sys
from pptx import Presentation
from pptx.util import Emu

BODY_MIN, FOOTER_MIN = 18, 12


def shapes(container):
    for sh in container.shapes:
        if sh.shape_type == 6:  # group
            yield from shapes(sh)
        else:
            yield sh


def overflows(sh):
    # ponytail: rough character-area estimate, not real text layout; the rendered contact sheet is the real check
    if not sh.width or not sh.height:
        return False
    need = 0.0
    for p in sh.text_frame.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        chars_per_line = max(1, Emu(sh.width).pt / (size * 0.52))
        lines = max(1, -(-len(p.text) // chars_per_line))
        need += lines * size * 1.2
    return need > Emu(sh.height).pt * 1.05


def check(path, max_words):
    errors, warnings = [], []
    for n, slide in enumerate(Presentation(path).slides, 1):
        words, has_title = 0, False
        for sh in shapes(slide):
            name = sh.name.lower()
            if name.startswith("title") or (sh.is_placeholder and sh.placeholder_format.type in (1, 3)):
                has_title = True
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            words += len(sh.text_frame.text.split())
            floor = FOOTER_MIN if name.startswith("footer") else BODY_MIN
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    if r.font.size is None:
                        errors.append(f"slide {n} '{sh.name}': no explicit font size on \"{r.text[:30]}\"")
                    elif r.font.size.pt < floor:
                        errors.append(f"slide {n} '{sh.name}': {r.font.size.pt:g}pt < {floor}pt on \"{r.text[:30]}\"")
            if overflows(sh):
                warnings.append(f"slide {n} '{sh.name}': text probably overflows its box")
        if not has_title:
            errors.append(f"slide {n}: no title shape")
        if words > max_words:
            warnings.append(f"slide {n}: {words} words (over {max_words}), crowded?")
    return errors, warnings


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("deck"); ap.add_argument("--max-words", type=int, default=90)
    a = ap.parse_args()
    errors, warnings = check(a.deck, a.max_words)
    for e in errors: print("FAIL", e)
    for w in warnings: print("WARN", w)
    print(f"{len(errors)} failures, {len(warnings)} warnings")
    sys.exit(1 if errors else 0)
